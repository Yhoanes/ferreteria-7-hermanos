from datetime import date
from decimal import Decimal

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from django.contrib import messages
from django.db.models import Count, DecimalField, ExpressionWrapper, F, Sum
from django.http import HttpResponse
from django.shortcuts import render
from django.utils.dateparse import parse_date

from administracion.decorators import requiere_permiso
from producto.models import Producto
from venta.models import DetalleVenta, Pago, Venta


TIPOS_REPORTE_VALIDOS = {
    "ventas",
    "inventario",
    "top_productos",
    "vendedores",
    "pagos",
}


def _money(valor):
    if valor is None:
        return Decimal("0.00")
    return Decimal(valor).quantize(Decimal("0.01"))


def _ventas_filtradas(request):
    fecha_inicio_raw = request.GET.get("desde", "").strip()
    fecha_fin_raw = request.GET.get("hasta", "").strip()

    fecha_inicio = parse_date(fecha_inicio_raw) if fecha_inicio_raw else None
    fecha_fin = parse_date(fecha_fin_raw) if fecha_fin_raw else None

    ventas = (
        Venta.objects
        .select_related("cliente", "usuario", "descuento")
        .order_by("-fecha", "-id")
    )

    if fecha_inicio_raw and not fecha_inicio:
        messages.error(request, "La fecha inicial no tiene un formato válido.")

    if fecha_fin_raw and not fecha_fin:
        messages.error(request, "La fecha final no tiene un formato válido.")

    if fecha_inicio and fecha_fin:
        if fecha_inicio > fecha_fin:
            messages.error(request, "La fecha inicial no puede ser posterior a la fecha final.")
            ventas = ventas.filter(fecha__date=date.today())
        else:
            ventas = ventas.filter(fecha__date__range=[fecha_inicio, fecha_fin])
    elif fecha_inicio:
        ventas = ventas.filter(fecha__date__gte=fecha_inicio)
    elif fecha_fin:
        ventas = ventas.filter(fecha__date__lte=fecha_fin)
    else:
        ventas = ventas.filter(fecha__date=date.today())

    return ventas, fecha_inicio_raw, fecha_fin_raw


def _aplicar_estilo_excel(ws):
    header_fill = PatternFill("solid", fgColor="1F2937")
    header_font = Font(color="FFFFFF", bold=True)
    border_color = "D9DEE3"
    thin_border = Border(
        left=Side(style="thin", color=border_color),
        right=Side(style="thin", color=border_color),
        top=Side(style="thin", color=border_color),
        bottom=Side(style="thin", color=border_color),
    )

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")

    for column, width in {
        "A": 15,
        "B": 42,
        "C": 28,
        "D": 22,
        "E": 22,
        "F": 18,
        "G": 18,
    }.items():
        ws.column_dimensions[column].width = width

    ws.freeze_panes = "A2"


def _exportar_excel(tipo_reporte, ventas):
    wb = openpyxl.Workbook()
    ws = wb.active

    if tipo_reporte == "inventario":
        ws.title = "Inventario"
        ws.append([
            "ID",
            "Código",
            "Producto",
            "Categoría",
            "Stock Actual",
            "Stock Mínimo",
            "Precio Venta (Bs)",
            "Valor Stock (Bs)",
            "Estado",
        ])

        productos = (
            Producto.objects
            .select_related("categoria", "marca")
            .all()
            .order_by("nombre", "id")
        )

        for producto in productos:
            categoria_nombre = producto.categoria.nombre if producto.categoria else "Sin Categoría"
            valor_stock = _money(producto.stock_actual * producto.precio_venta)

            if producto.stock_actual <= 0:
                estado = "SIN STOCK"
            elif producto.stock_actual <= producto.stock_minimo:
                estado = "CRÍTICO"
            else:
                estado = "OK"

            ws.append([
                producto.id,
                producto.codigo,
                producto.nombre,
                categoria_nombre,
                producto.stock_actual,
                producto.stock_minimo,
                float(producto.precio_venta),
                float(valor_stock),
                estado,
            ])

        nombre_archivo = "Reporte_Inventario_Ferreteria.xlsx"

    elif tipo_reporte == "top_productos":
        ws.title = "Top Productos"
        ws.append([
            "Código",
            "Producto",
            "Unidades Vendidas",
            "Total Recaudado (Bs)",
        ])

        top_productos_data = (
            DetalleVenta.objects
            .filter(venta__in=ventas)
            .values("producto__codigo", "producto__nombre")
            .annotate(
                total_vendido=Sum("cantidad"),
                total_recaudado=Sum("subtotal"),
            )
            .order_by("-total_vendido", "-total_recaudado")
        )

        for item in top_productos_data:
            ws.append([
                item["producto__codigo"] or "Sin código",
                item["producto__nombre"] or "Producto no disponible",
                item["total_vendido"] or 0,
                float(_money(item["total_recaudado"])),
            ])

        nombre_archivo = "Reporte_Top_Productos_Ferreteria.xlsx"

    elif tipo_reporte == "vendedores":
        ws.title = "Vendedores"
        ws.append([
            "Vendedor",
            "Transacciones Realizadas",
            "Total Facturado (Bs)",
        ])

        vendedores_data = (
            ventas
            .values("usuario__nombre")
            .annotate(
                c_ventas=Count("id"),
                t_facturado=Sum("total"),
            )
            .order_by("-t_facturado")
        )

        for item in vendedores_data:
            ws.append([
                item["usuario__nombre"] or "Sistema / Virtual",
                item["c_ventas"] or 0,
                float(_money(item["t_facturado"])),
            ])

        nombre_archivo = "Reporte_Rendimiento_Vendedores.xlsx"

    elif tipo_reporte == "pagos":
        ws.title = "Pagos"
        ws.append([
            "Método de Pago",
            "Transacciones Realizadas",
            "Total Recaudado (Bs)",
        ])

        pagos_data = (
            Pago.objects
            .filter(venta__in=ventas)
            .values("tipo_pago__nombre")
            .annotate(
                cantidad_operaciones=Count("id"),
                total_recaudado=Sum("monto"),
            )
            .order_by("-total_recaudado")
        )

        for item in pagos_data:
            ws.append([
                item["tipo_pago__nombre"] or "No especificado",
                item["cantidad_operaciones"] or 0,
                float(_money(item["total_recaudado"])),
            ])

        nombre_archivo = "Reporte_Pagos_Caja_Ferreteria.xlsx"

    else:
        ws.title = "Ventas"
        ws.append([
            "ID Venta",
            "Fecha",
            "Cliente",
            "Documento Cliente",
            "Vendedor",
            "Descuento",
            "Total (Bs)",
        ])

        for venta in ventas:
            fecha_limpia = venta.fecha.replace(tzinfo=None)

            ws.append([
                venta.id,
                fecha_limpia,
                venta.cliente.nombre if venta.cliente else "Sin cliente",
                venta.cliente.documento_identidad if venta.cliente else "Sin documento",
                venta.usuario.nombre if venta.usuario else "Sistema / Virtual",
                venta.descuento.nombre if venta.descuento else "Sin descuento",
                float(_money(venta.total)),
            ])

        nombre_archivo = "Reporte_Ventas_Ferreteria.xlsx"

    _aplicar_estilo_excel(ws)

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{nombre_archivo}"'
    wb.save(response)

    return response


@requiere_permiso("view_venta")
def reporte_ventas(request):
    ventas, fecha_inicio, fecha_fin = _ventas_filtradas(request)
    tipo_reporte = request.GET.get("tipo", "ventas").strip()

    if tipo_reporte not in TIPOS_REPORTE_VALIDOS:
        messages.error(request, "El tipo de reporte solicitado no es válido.")
        tipo_reporte = "ventas"

    total_dinero = _money(ventas.aggregate(total=Sum("total"))["total"])
    cantidad_ventas = ventas.count()

    if "exportar_excel" in request.GET:
        return _exportar_excel(tipo_reporte, ventas)

    productos = None
    top_productos = None
    vendedores = None
    resumen_pagos = None

    valor_inventario = Decimal("0.00")
    productos_alerta = 0
    productos_sin_stock = 0
    total_productos = 0
    mejor_vendedor = "N/A"
    total_vendedores = 0
    total_metodos = 0

    if tipo_reporte == "inventario":
        productos = (
            Producto.objects
            .select_related("categoria", "marca")
            .all()
            .order_by("nombre", "id")
        )

        valor_stock_expression = ExpressionWrapper(
            F("stock_actual") * F("precio_venta"),
            output_field=DecimalField(max_digits=14, decimal_places=2),
        )

        valor_inventario = _money(
            Producto.objects.aggregate(
                capital_total=Sum(valor_stock_expression)
            )["capital_total"]
        )

        productos_alerta = Producto.objects.filter(
            stock_actual__lte=F("stock_minimo")
        ).count()

        productos_sin_stock = Producto.objects.filter(
            stock_actual__lte=0
        ).count()

        total_productos = Producto.objects.count()

    elif tipo_reporte == "top_productos":
        top_productos = (
            DetalleVenta.objects
            .filter(venta__in=ventas)
            .values("producto__codigo", "producto__nombre")
            .annotate(
                total_vendido=Sum("cantidad"),
                total_recaudado=Sum("subtotal"),
            )
            .order_by("-total_vendido", "-total_recaudado")
        )

    elif tipo_reporte == "vendedores":
        vendedores = (
            ventas
            .values("usuario__nombre")
            .annotate(
                c_ventas=Count("id"),
                t_facturado=Sum("total"),
            )
            .order_by("-t_facturado")
        )

        total_vendedores = vendedores.count()

        if total_vendedores > 0:
            mejor_vendedor = vendedores[0]["usuario__nombre"] or "Sistema Virtual"

    elif tipo_reporte == "pagos":
        resumen_pagos = (
            Pago.objects
            .filter(venta__in=ventas)
            .values("tipo_pago__nombre")
            .annotate(
                cantidad_operaciones=Count("id"),
                total_recaudado=Sum("monto"),
            )
            .order_by("-total_recaudado")
        )

        total_metodos = resumen_pagos.count()

    return render(
        request,
        "reporte_ventas.html",
        {
            "ventas": ventas,
            "productos": productos,
            "top_productos": top_productos,
            "vendedores": vendedores,
            "resumen_pagos": resumen_pagos,
            "total_dinero": total_dinero,
            "cantidad_ventas": cantidad_ventas,
            "valor_inventario": valor_inventario,
            "productos_alerta": productos_alerta,
            "productos_sin_stock": productos_sin_stock,
            "total_productos": total_productos,
            "mejor_vendedor": mejor_vendedor,
            "total_vendedores": total_vendedores,
            "total_metodos": total_metodos,
            "desde": fecha_inicio,
            "hasta": fecha_fin,
            "tipo": tipo_reporte,
            "hoy": date.today().strftime("%Y-%m-%d"),
        },
    )